"""Synthetic fixtures only: no real course documents or provider calls."""

import importlib
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile


SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "cadence" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import extract


def available(name):
    return importlib.util.find_spec(name) is not None


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def file(self, name, value):
        path = self.directory / name
        path.write_bytes(value)
        return path

    def test_utf8_markdown_line_numbers_and_nuls(self):
        path = self.file("lesson.MD", "\ufeff# Purpose\n\nCafé\x00 activity\nSecond line\n".encode("utf-8"))
        result = extract.extract_file(path)
        self.assertEqual(result["units"], [
            {"text": "# Purpose", "locator": "lines 1–1"},
            {"text": "Café activity\nSecond line", "locator": "lines 3–4"},
        ])
        self.assertIn("not fetched", result["warnings"][0])

    def test_plain_text_does_not_import_optional_packages(self):
        path = self.file("lesson.txt", b"Read this lesson.")
        with patch.object(extract, "_dependency", side_effect=AssertionError("unexpected import")):
            self.assertEqual(extract.extract_file(path)["units"][0]["text"], "Read this lesson.")

    def test_invalid_text_encoding(self):
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            extract.extract_file(self.file("lesson.txt", b"\xff\xfeinvalid"))

    def test_empty_and_unsupported_input(self):
        with self.assertRaisesRegex(ValueError, "No readable text"):
            extract.extract_file(self.file("empty.txt", b"\x00 \n\n"))
        with self.assertRaisesRegex(ValueError, "Supported source formats"):
            extract.extract_file(self.file("course.mbz", b"anything"))
        with self.assertRaisesRegex(ValueError, "does not exist"):
            extract.extract_file(self.directory / "absent.txt")

    def test_file_and_text_limits(self):
        path = self.file("large.txt", b"123456789")
        with patch.object(extract, "MAX_FILE_BYTES", 8):
            with self.assertRaisesRegex(ValueError, "1 GB"):
                extract.extract_file(path)
        with patch.object(extract, "MAX_TEXT_CHARS", 8):
            with self.assertRaisesRegex(ValueError, "8 characters"):
                extract.extract_file(path)

    def test_missing_optional_dependency_is_actionable(self):
        path = self.file("source.pdf", b"placeholder")
        with patch.object(importlib, "import_module", side_effect=ModuleNotFoundError("pypdf")):
            with self.assertRaisesRegex(ValueError, "requires pypdf.*requirements.txt"):
                extract.extract_file(path)

    def fake_pdf(self, texts, encrypted=False):
        class Page:
            def __init__(self, value):
                self.value = value

            def extract_text(self):
                if isinstance(self.value, Exception):
                    raise self.value
                return self.value

        class Reader:
            is_encrypted = encrypted
            pages = [Page(text) for text in texts]

            def close(self):
                pass

        class PdfModule:
            @staticmethod
            def PdfReader(path):
                return Reader()

        return PdfModule

    def test_mixed_pdf_preserves_page_numbers_and_lists_skips(self):
        path = self.file("mixed.pdf", b"mock")
        fake = self.fake_pdf(["Read\x00able", "", None, "Last page", "\x00  "])
        with patch.object(extract, "_dependency", return_value=fake):
            result = extract.extract_file(path)
        self.assertEqual(result["units"], [
            {"text": "Readable", "locator": "page 1", "page": 1},
            {"text": "Last page", "locator": "page 4", "page": 4},
        ])
        self.assertTrue(any("Pages 2–3, 5" in warning for warning in result["warnings"]))

    def test_pdf_with_one_failed_page_keeps_other_pages(self):
        fake = self.fake_pdf([RuntimeError("bad page"), "good page"])
        with patch.object(extract, "_dependency", return_value=fake):
            result = extract.extract_file(self.file("mixed.pdf", b"mock"))
        self.assertEqual(result["units"][0]["page"], 2)
        self.assertTrue(any("Pages 1 could not be extracted" in warning for warning in result["warnings"]))

    def test_pdf_scanned_encrypted_and_page_limit_errors(self):
        path = self.file("source.pdf", b"mock")
        for fake, message in [
            (self.fake_pdf(["", None]), "Run OCR externally"),
            (self.fake_pdf(["a"], encrypted=True), "Encrypted PDFs"),
            (self.fake_pdf(["a"] * 1001), "1,000 pages"),
        ]:
            with self.subTest(message=message), patch.object(extract, "_dependency", return_value=fake):
                with self.assertRaisesRegex(ValueError, message):
                    extract.extract_file(path)

    @unittest.skipUnless(available("pypdf"), "pypdf is optional")
    def test_real_blank_pdf_and_corrupt_pdf(self):
        from pypdf import PdfWriter

        path = self.directory / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        with path.open("wb") as stream:
            writer.write(stream)
        with self.assertRaisesRegex(ValueError, "Run OCR externally"):
            extract.extract_file(path)
        with self.assertRaisesRegex(ValueError, "Could not read"):
            extract.extract_file(self.file("corrupt.pdf", b"not a pdf"))

    @unittest.skipUnless(available("docx"), "python-docx is optional")
    def test_docx_heading_paragraph_and_nested_table(self):
        from docx import Document

        document = Document()
        document.add_heading("Lesson purpose", level=1)
        document.add_paragraph("Practice making a decision.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Activity"
        table.cell(0, 1).text = "Objective"
        table.cell(1, 0).text = "Discuss"
        table.cell(1, 1).text = "Explain"
        nested = table.cell(1, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested guidance"
        path = self.directory / "lesson.docx"
        document.save(path)
        result = extract.extract_file(path)
        self.assertEqual(result["units"][0], {"text": "Lesson purpose", "locator": "paragraph 1 (Heading 1)"})
        self.assertEqual(result["units"][1]["locator"], "paragraph 2")
        self.assertEqual(result["units"][2], {"text": "Activity\tObjective", "locator": "table 1, row 1"})
        self.assertIn("Nested guidance", result["units"][3]["text"])

    @unittest.skipUnless(available("pptx"), "python-pptx is optional")
    def test_pptx_titles_tables_notes_groups_and_original_slide_numbers(self):
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Decision making"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
        table.cell(0, 0).text = "Activity"
        table.cell(0, 1).text = "Time"
        table.cell(1, 0).text = "Discussion"
        table.cell(1, 1).text = "10 minutes"
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text = "Grouped text"
        slide.notes_slide.notes_text_frame.text = "Ask learners to justify their choices."
        presentation.slides.add_slide(presentation.slide_layouts[6])
        last = presentation.slides.add_slide(presentation.slide_layouts[5])
        last.shapes.title.text = "Debrief"
        path = self.directory / "lesson.pptx"
        presentation.save(path)
        result = extract.extract_file(path)
        texts = [unit["text"] for unit in result["units"]]
        self.assertIn("Activity\tTime", texts)
        self.assertIn("Discussion\t10 minutes", texts)
        self.assertIn("Grouped text", texts)
        notes = next(unit for unit in result["units"] if unit["locator"] == "slide 1, speaker notes")
        self.assertIn("justify", notes["text"])
        self.assertEqual(result["units"][-1]["slide"], 3)
        self.assertTrue(any("Slides 2" in warning for warning in result["warnings"]))

    def package(self, entries):
        path = self.directory / "package.docx"
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in entries:
                archive.writestr(name, data)
        return path

    def test_unsafe_archive_members_are_rejected(self):
        for name in ["../outside.xml", "/absolute.xml", "C:/windows.xml", "word\\document.xml"]:
            with self.subTest(name=name):
                path = self.package([(name, b"a")])
                with self.assertRaisesRegex(ValueError, "unsafe archive path"):
                    extract._check_package(path, "word/document.xml")

    def test_archive_symlinks_and_expanded_limit(self):
        link = zipfile.ZipInfo("word/document.xml")
        link.create_system = 3
        link.external_attr = 0o120777 << 16
        with self.assertRaisesRegex(ValueError, "Links are not supported"):
            extract._check_package(self.package([(link, b"target")]), "word/document.xml")
        path = self.package([("word/document.xml", b"123456789")])
        with patch.object(extract, "MAX_EXPANDED_BYTES", 8):
            with self.assertRaisesRegex(ValueError, "Expanded.*1 GB"):
                extract._check_package(path, "word/document.xml")

    def test_entity_declarations_are_rejected(self):
        for text in [
            '<!DOCTYPE doc [<!ENTITY x SYSTEM "file:///not-read">]><doc>&x;</doc>',
            " " * 65530 + "<!DOCTYPE doc><doc/>",
        ]:
            for encoding in ["utf-8", "utf-16"]:
                with self.subTest(encoding=encoding):
                    path = self.package([("word/document.xml", text.encode(encoding))])
                    with self.assertRaisesRegex(ValueError, "DTD and entity"):
                        extract._check_package(path, "word/document.xml")

    def test_duplicate_archive_paths_are_rejected(self):
        path = self.package([("word/document.xml", b"one"), ("word/./document.xml", b"two")])
        with self.assertRaisesRegex(ValueError, "duplicate archive paths"):
            extract._check_package(path, "word/document.xml")


if __name__ == "__main__":
    unittest.main()
