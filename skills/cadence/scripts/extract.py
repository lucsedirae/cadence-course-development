"""Local, citation-preserving text extraction for Cadence's two libraries.

The core uses only the standard library. Binary-format parsers are imported only
when their format is requested. Nothing is executed, fetched, or OCR'd.
"""

from __future__ import annotations

import importlib
from pathlib import Path, PurePosixPath
import stat
import zipfile


MAX_FILE_BYTES = 1024 * 1024 * 1024
MAX_EXPANDED_BYTES = 1024 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10_000
MAX_TEXT_CHARS = 64 * 1024 * 1024
MAX_PDF_PAGES = 1000
SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".docx", ".pdf", ".pptx"}


def _dependency(module: str, package: str):
    try:
        return importlib.import_module(module)
    except ImportError as exc:
        raise ValueError(
            f"Reading this format requires {package}. Install the plugin's "
            "requirements.txt in your Python environment."
        ) from exc


def _ranges(numbers: list[int]) -> str:
    groups = []
    for number in sorted(set(numbers)):
        if groups and number == groups[-1][1] + 1:
            groups[-1][1] = number
        else:
            groups.append([number, number])
    return ", ".join(str(a) if a == b else f"{a}–{b}" for a, b in groups)


class _Result:
    def __init__(self):
        self.units: list[dict] = []
        self.warnings: list[str] = []
        self.characters = 0

    def add(self, text: str, locator: str, **location):
        text = text.replace("\x00", "").strip()
        if not text:
            return
        self.characters += len(text)
        if self.characters > MAX_TEXT_CHARS:
            raise ValueError(
                f"Extracted text must be {MAX_TEXT_CHARS:,} characters or fewer. "
                "Split this document into smaller source files."
            )
        self.units.append({"text": text, "locator": locator, **location})

    def finish(self):
        if not self.units:
            raise ValueError("No readable text was found in this file.")
        return {"units": self.units, "warnings": self.warnings}


def _check_package(path: Path, required_part: str):
    """Inspect OOXML in place; never extract archive members to the filesystem."""
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("Office documents may contain at most 10,000 archive entries.")
            if sum(entry.file_size for entry in members) > MAX_EXPANDED_BYTES:
                raise ValueError("Expanded Word and PowerPoint documents must be 1 GB or smaller.")
            seen = set()
            for entry in members:
                raw = entry.filename
                member = PurePosixPath(raw)
                if (
                    member.is_absolute() or ".." in member.parts or "\\" in raw
                    or ":" in raw or "\x00" in raw or not member.parts
                ):
                    raise ValueError("The Office document contains an unsafe archive path.")
                name = str(member)
                if name in seen:
                    raise ValueError("The Office document contains duplicate archive paths.")
                seen.add(name)
                if stat.S_ISLNK(entry.external_attr >> 16):
                    raise ValueError("Links are not supported inside Office document archives.")
                if entry.flag_bits & 1:
                    raise ValueError("Encrypted Office document archives are not supported.")
                # Office files have no need for DTDs/entities. Check UTF-8 and
                # UTF-16/32 declarations without asking an XML parser to resolve
                # them. A carry buffer catches declarations across chunk edges.
                if name.lower().endswith((".xml", ".rels")) and not entry.is_dir():
                    with archive.open(entry) as stream:
                        carry = b""
                        while chunk := stream.read(64 * 1024):
                            probe = carry + chunk.replace(b"\x00", b"").lower()
                            if b"<!doctype" in probe or b"<!entity" in probe:
                                raise ValueError("DTD and entity declarations are not supported in Office documents.")
                            carry = probe[-32:]
            if required_part not in seen or "[Content_Types].xml" not in seen:
                raise ValueError("This file is not a valid document of its stated Office format.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Could not read this Office document. Save a valid .docx or .pptx file and try again.") from exc


def _text(path: Path) -> _Result:
    result = _Result()
    with path.open("rb") as source:
        raw = source.read(MAX_TEXT_CHARS * 4 + 4)
    if len(raw) > MAX_TEXT_CHARS * 4 + 3:
        raise ValueError(f"Extracted text must be {MAX_TEXT_CHARS:,} characters or fewer.")
    try:
        text = raw.decode("utf-8-sig").replace("\x00", "")
    except UnicodeDecodeError as exc:
        raise ValueError("Text and Markdown files must use UTF-8 encoding. Save this file as UTF-8 and try again.") from exc
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError(f"Extracted text must be {MAX_TEXT_CHARS:,} characters or fewer.")
    lines = text.splitlines()
    start, paragraph = 1, []
    for number, line in enumerate(lines, 1):
        if line.strip():
            if not paragraph:
                start = number
            paragraph.append(line)
        elif paragraph:
            result.add("\n".join(paragraph), f"lines {start}–{number - 1}")
            paragraph = []
    if paragraph:
        result.add("\n".join(paragraph), f"lines {start}–{len(lines)}")
    if path.suffix.lower() in {".md", ".markdown"}:
        result.warnings.append("Markdown source text is indexed as written; linked files, images, and websites are not fetched.")
    return result


def _pdf(path: Path) -> _Result:
    pypdf = _dependency("pypdf", "pypdf")
    result = _Result()
    reader = pypdf.PdfReader(str(path))
    try:
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported. Save an unencrypted copy and try again.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"PDFs must have {MAX_PDF_PAGES:,} pages or fewer.")
        blank, failed = [], []
        for number, page in enumerate(reader.pages, 1):
            try:
                value = (page.extract_text() or "").replace("\x00", "").strip()
            except Exception:
                failed.append(number)
                continue
            if value:
                result.add(value, f"page {number}", page=number)
            else:
                blank.append(number)
        if not result.units:
            if failed:
                raise ValueError(
                    f"No readable text was found. PDF pages {_ranges(failed)} could not be extracted. "
                    "Try re-exporting the PDF; scanned pages require external OCR."
                )
            raise ValueError("No selectable text was found in this PDF. Run OCR externally and upload the resulting PDF or text file.")
        result.warnings.append("PDF text is extracted by page. Check tables and reading order against the original; images are not analyzed.")
        if blank:
            result.warnings.append(f"Pages {_ranges(blank)} had no selectable text and were skipped. OCR was not performed.")
        if failed:
            result.warnings.append(f"Pages {_ranges(failed)} could not be extracted and were skipped.")
    finally:
        close = getattr(reader, "close", None)
        if close:
            close()
    return result


def _docx(path: Path) -> _Result:
    docx = _dependency("docx", "python-docx")
    _check_package(path, "word/document.xml")
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    result = _Result()
    document = docx.Document(str(path))

    def table_rows(table):
        for row in table.rows:
            seen, texts = set(), []
            for cell in row.cells:
                if cell._tc in seen:
                    texts.append("")
                    continue
                seen.add(cell._tc)
                values = []
                for child in cell.iter_inner_content():
                    if isinstance(child, Table):
                        values.extend(table_rows(child))
                    else:
                        values.append(child.text)
                texts.append("\n".join(values))
            yield "\t".join(texts)

    paragraph_number = table_number = 0
    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraph_number += 1
            paragraph = Paragraph(element, document)
            style = paragraph.style.name if paragraph.style else ""
            locator = f"paragraph {paragraph_number}"
            if style.startswith("Heading") or style == "Title":
                locator += f" ({style})"
            result.add(paragraph.text, locator)
        elif element.tag == qn("w:tbl"):
            table_number += 1
            for row_number, value in enumerate(table_rows(Table(element, document)), 1):
                result.add(value, f"table {table_number}, row {row_number}")
    result.warnings.append(
        "Word body headings, paragraphs, and tables were imported. Headers, footers, "
        "footnotes, text boxes, images, and tracked-change details are not analyzed; "
        "check complex layouts against the original."
    )
    return result


def _pptx(path: Path) -> _Result:
    pptx = _dependency("pptx", "python-pptx")
    _check_package(path, "ppt/presentation.xml")
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    result = _Result()
    presentation = pptx.Presentation(str(path))

    def shapes(collection, prefix=""):
        for index, shape in enumerate(collection, 1):
            position = f"{prefix}.{index}" if prefix else str(index)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes(shape.shapes, position)
            else:
                yield position, shape

    empty_slides = []
    for number, slide in enumerate(presentation.slides, 1):
        before = len(result.units)
        title = slide.shapes.title
        for position, shape in shapes(slide.shapes):
            locator = f"slide {number}, shape {position}"
            if title is not None and shape.shape_id == title.shape_id:
                locator = f"slide {number}, title"
            if shape.has_table:
                for row_number, row in enumerate(shape.table.rows, 1):
                    value = "\t".join("" if cell.is_spanned else cell.text for cell in row.cells)
                    result.add(value, f"{locator}, table row {row_number}", slide=number)
            elif shape.has_text_frame:
                result.add(shape.text_frame.text, locator, slide=number)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                result.add(frame.text, f"slide {number}, speaker notes", slide=number)
        if len(result.units) == before:
            empty_slides.append(number)
    result.warnings.append(
        "PowerPoint text, tables, and speaker notes were imported in slide order. "
        "Images, charts, diagrams, animations, media, and visual layout are not analyzed."
    )
    if empty_slides:
        result.warnings.append(f"Slides {_ranges(empty_slides)} had no readable text or speaker notes and were skipped.")
    return result


def extract_file(path: str | Path) -> dict:
    """Return {units: [{text, locator, page?/slide?}], warnings: [str]}.

    Locators always refer to the original source. Units contain extracted text,
    never generated interpretation. Invalid/unsupported input raises ValueError.
    """
    source = Path(path).expanduser()
    if not source.is_file():
        raise ValueError(f"Source file does not exist or is not a regular file: {source}")
    try:
        if source.stat().st_size > MAX_FILE_BYTES:
            raise ValueError("Files must be 1 GB or smaller.")
        suffix = source.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(
                "Supported source formats are .txt, .md, .markdown, .docx, .pdf, and .pptx. "
                "Export other formats to one of these before importing."
            )
        parser = {".docx": _docx, ".pdf": _pdf, ".pptx": _pptx}.get(suffix, _text)
        return parser(source).finish()
    except ValueError:
        raise
    except PermissionError as exc:
        raise ValueError(f"Cannot read source file: {source.name}. Check its file permissions.") from exc
    except Exception as exc:
        raise ValueError(
            f"Could not read {source.name}. Check that the file is complete and its contents match its extension."
        ) from exc
